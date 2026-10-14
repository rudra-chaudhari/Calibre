#!/usr/bin/env python3
"""
Universal Timetable Excel (.xlsx) -> JSON converter
===================================================





import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

# ----------------------------------------------------------------------------
# Regexes / constants
# ----------------------------------------------------------------------------
TIME_RANGE_RE = re.compile(
    r"(\d{1,2})[:.](\d{2})(?:[:.]\d{2})?\s*([AaPp][Mm])?\s*(?:-|–|—|to|TO)\s*"
    r"(\d{1,2})[:.](\d{2})(?:[:.]\d{2})?\s*([AaPp][Mm])?"
)

DAY_RE = re.compile(
    r"^\W*(?P<no>\d+)?\W*\s*(?P<day>mon|tue|wed|thu|fri|sat|sun)[a-z]*\b", re.I
)
DAYS = {"mon": (1, "Mon", "Monday"), "tue": (2, "Tue", "Tuesday"),
        "wed": (3, "Wed", "Wednesday"), "thu": (4, "Thu", "Thursday"),
        "fri": (5, "Fri", "Friday"), "sat": (6, "Sat", "Saturday"),
        "sun": (7, "Sun", "Sunday")}

# "EE 229(S2 )", "CS6007", "HS 101", "ES 6001", "MA 105 Calculus"
COURSE_RE = re.compile(
    r"^(?P<dept>[A-Z]{1,6})\s*-?\s*(?P<num>\d{2,5}[A-Z]{0,2}(?:\.\d+)?)"
    r"\s*(?:\(\s*(?P<sec>[^)]*?)\s*\))?\s*(?P<title>.*)$"
)
ROOM_BRACKET_RE = re.compile(r"^\[(?P<room>.*)\]$")
ROOM_PREFIX_RE = re.compile(r"^(?:class\s*room|classroom|room|venue|hall)\s*[:\-]\s*(?P<room>.+)$", re.I)


def clean(s):
    """Normalise whitespace / weird chars."""
    s = str(s).replace("\xa0", " ").replace("\r", "")
    return s


def squeeze(s):
    return re.sub(r"\s+", " ", s).strip()


# ----------------------------------------------------------------------------
# Time & day parsing
# ----------------------------------------------------------------------------
def _to_24h(h, m, ampm):
    h, m = int(h), int(m)
    if ampm:
        ampm = ampm.lower()
        if ampm == "pm" and h < 12:
            h += 12
        if ampm == "am" and h == 12:
            h = 0
    return h, m


def parse_time_range(value):
    """'08:00:00- 09:25:00' -> ('08:00','09:25') else None"""
    if value is None:
        return None
    m = TIME_RANGE_RE.search(clean(value))
    if not m:
        return None
    h1, m1, ap1, h2, m2, ap2 = m.groups()
    ap1 = ap1 or ap2  # "8.00-9.25 AM" -> dono AM
    sh, sm = _to_24h(h1, m1, ap1)
    eh, em = _to_24h(h2, m2, ap2)
    if not (0 <= sh < 24 and 0 <= eh < 24 and sm < 60 and em < 60):
        return None
    return f"{sh:02d}:{sm:02d}", f"{eh:02d}:{em:02d}"


def parse_day(value):
    """'1. Mon' / 'Monday' -> (day_no, 'Mon', 'Monday') else None"""
    if value is None:
        return None
    text = clean(value).strip()
    if len(text) > 25:  # lamba text day label nahi ho sakta
        return None
    m = DAY_RE.match(text)
    if not m:
        return None
    key = m.group("day").lower()
    no, short, full = DAYS[key]
    return (int(m.group("no")) if m.group("no") else no), short, full


# ----------------------------------------------------------------------------
# Cell parsing  (ek cell -> multiple courses)
# ----------------------------------------------------------------------------
def parse_cell(value):
    """
    Cell text -> list of course dicts.
    Naya course tab shuru hota hai jab koi line course-code jaisi dikhe
    (blank line ki zaroorat nahi).
    """
    if value is None:
        return []
    text = clean(value).strip()
    if not text:
        return []

    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    blocks, current = [], None

    for ln in lines:
        if ROOM_BRACKET_RE.match(ln) or ROOM_PREFIX_RE.match(ln):
            if current is None:
                current = {"header": None, "room": None, "people": [], "extra": []}
                blocks.append(current)
            m = ROOM_BRACKET_RE.match(ln) or ROOM_PREFIX_RE.match(ln)
            room = m.group("room").strip()
            room = ROOM_PREFIX_RE.sub(lambda mm: mm.group("room"), room).strip()
            current["room"] = squeeze(room)
            continue

        cm = COURSE_RE.match(ln)
        if cm:
            current = {"header": cm, "room": None, "people": [], "extra": []}
            blocks.append(current)
        else:
            if current is None:
                current = {"header": None, "room": None, "people": [], "extra": []}
                blocks.append(current)
            current["people"].append(squeeze(ln))

    out = []
    for b in blocks:
        h = b["header"]
        if h:
            title = squeeze(h.group("title") or "")
            out.append({
                "department": h.group("dept").upper(),
                "course_code": f"{h.group('dept').upper()} {h.group('num')}",
                "section": squeeze(h.group("sec")) if h.group("sec") and h.group("sec").strip() else None,
                "title": title or None,
                "room": b["room"],
                "instructor": ", ".join(b["people"]) or None,
            })
        else:  # header nahi mila -> raw text bachao, data lose nahi hoga
            out.append({
                "department": None,
                "course_code": None,
                "section": None,
                "title": None,
                "room": b["room"],
                "instructor": ", ".join(b["people"]) or None,
                "raw": squeeze(" | ".join(b["people"])) or None,
            })
    return out


# ----------------------------------------------------------------------------
# Sheet -> grid (merged cells ke saath) + layout detection
# ----------------------------------------------------------------------------
def sheet_to_grid(ws):
    nrows, ncols = ws.max_row, ws.max_column
    grid = [[ws.cell(r, c).value for c in range(1, ncols + 1)] for r in range(1, nrows + 1)]
    merges = []  # (r1, c1, r2, c2) 0-indexed
    for rng in ws.merged_cells.ranges:
        r1, c1, r2, c2 = rng.min_row - 1, rng.min_col - 1, rng.max_row - 1, rng.max_col - 1
        merges.append((r1, c1, r2, c2))
    return grid, merges


def transpose(grid, merges):
    if not grid:
        return grid, merges
    t = [list(row) for row in zip(*grid)]
    tm = [(c1, r1, c2, r2) for (r1, c1, r2, c2) in merges]
    return t, tm


def find_time_header(grid, max_scan=15):
    """Row jisme >=2 time-ranges ho -> (row_idx, {col_idx: (start,end)})"""
    best = None
    for r, row in enumerate(grid[:max_scan]):
        cols = {}
        for c, v in enumerate(row):
            t = parse_time_range(v)
            if t:
                cols[c] = t
        if len(cols) >= 2 and (best is None or len(cols) > len(best[1])):
            best = (r, cols)
    return best


def detect_layout(grid, merges):
    """Time header row mile to normal; warna transpose karke try karo."""
    hdr = find_time_header(grid)
    if hdr:
        return grid, merges, hdr
    tg, tm = transpose(grid, merges)
    hdr = find_time_header(tg)
    if hdr:
        return tg, tm, hdr
    return grid, merges, None


def convert_sheet(ws):
    grid, merges = sheet_to_grid(ws)
    grid, merges, hdr = detect_layout(grid, merges)
    if not hdr:
        return None  # is sheet me time-table nahi mila

    hdr_row, slots = hdr
    ncols = max(len(r) for r in grid) if grid else 0

    # day label wala column dhundo: header ke neeche sabse zyada day-names jahan
    day_col, best = 0, -1
    for c in range(ncols):
        n = sum(1 for r in range(hdr_row + 1, len(grid)) if c < len(grid[r]) and parse_day(grid[r][c]))
        if n > best:
            day_col, best = c, n
    if best <= 0:
        return None

    # merged: anchor -> span, aur non-anchor cells skip
    anchor_span, covered = {}, set()
    for (r1, c1, r2, c2) in merges:
        anchor_span[(r1, c1)] = (r2, c2)
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                if (rr, cc) != (r1, c1):
                    covered.add((rr, cc))

    # merged day-label (vertical merge) support: pichla day aage carry karo
    entries, last_day = [], None
    slot_cols = sorted(slots)

    for r in range(hdr_row + 1, len(grid)):
        row = grid[r]
        day = parse_day(row[day_col]) if day_col < len(row) else None
        if day:
            last_day = day
        elif (r, day_col) in covered and last_day:
            day = last_day
        if not day:
            continue
        day_no, day_short, day_full = day

        for c in slot_cols:
            if (r, c) in covered or c >= len(row):
                continue
            courses = parse_cell(row[c])
            if not courses:
                continue
            start, end = slots[c]
            # horizontal merge -> end time last merged slot ka
            if (r, c) in anchor_span:
                _, c2 = anchor_span[(r, c)]
                last = max((cc for cc in slot_cols if c <= cc <= c2), default=c)
                end = slots[last][1]
            for course in courses:
                entries.append({
                    "day": day_short,
                    "day_name": day_full,
                    "day_no": day_no,
                    "start_time": start,
                    "end_time": end,
                    **course,
                })
    return entries


# ----------------------------------------------------------------------------
# Output shaping
# ----------------------------------------------------------------------------
def shape(entries, fmt):
    entries = sorted(entries, key=lambda e: (e["day_no"], e["start_time"], e.get("course_code") or ""))
    if fmt == "day":
        out = {}
        for e in entries:
            out.setdefault(e["day"], []).append({k: v for k, v in e.items() if k not in ("day", "day_name", "day_no")})
        return out
    if fmt == "course":
        out = {}
        for e in entries:
            key = e.get("course_code") or "UNKNOWN"
            out.setdefault(key, []).append({k: v for k, v in e.items() if k not in ("course_code", "department")})
        return out
    return entries


def process_file(path, fmt="flat", dept=None):
    wb = load_workbook(path, data_only=True)
    sheets, total = {}, 0
    for ws in wb.worksheets:
        entries = convert_sheet(ws)
        if entries is None:
            print(f"   [skip] sheet '{ws.title}': timetable structure nahi mila", file=sys.stderr)
            continue
        if dept:
            entries = [e for e in entries if (e.get("department") or "").upper() == dept.upper()]
        total += len(entries)
        sheets[ws.title] = entries
    return sheets, total


def build_result(path, sheets, fmt):
    shaped = {name: shape(ents, fmt) for name, ents in sheets.items()}
    res = {"source_file": Path(path).name, "format": fmt}
    if len(shaped) == 1:
        res["timetable"] = next(iter(shaped.values()))
    else:
        res["sheets"] = shaped
    return res


def collect_inputs(items):
    files = []
    for it in items:
        p = Path(it)
        if p.is_dir():
            files += sorted(x for x in p.glob("*.xlsx") if not x.name.startswith("~$"))
        elif p.exists():
            files.append(p)
        else:
            print(f"[warn] nahi mila: {it}", file=sys.stderr)
    return files


def main():
    ap = argparse.ArgumentParser(description="Universal timetable Excel -> JSON")
    ap.add_argument("inputs", nargs="+", help=".xlsx file(s) ya folder")
    ap.add_argument("-o", "--output", help="output path (sirf single input ke liye)")
    ap.add_argument("--merge", metavar="OUT.json", help="saari files ek combined JSON me")
    ap.add_argument("--format", choices=["flat", "day", "course"], default="flat")
    ap.add_argument("--dept", help="sirf is department ke courses (jaise CS, EE, MA)")
    args = ap.parse_args()

    files = collect_inputs(args.inputs)
    if not files:
        sys.exit("Koi input file nahi mili")
    if args.output and len(files) > 1:
        sys.exit("-o sirf ek file ke saath; multiple files ke liye --merge use karo")

    merged = []
    for f in files:
        sheets, total = process_file(f, args.format, args.dept)
        if args.merge:
            for ents in sheets.values():
                merged += ents
            print(f"[OK] {f.name}: {total} entries")
            continue
        out = Path(args.output) if args.output else f.with_suffix(".json")
        with open(out, "w", encoding="utf-8") as fh:
            json.dump(build_result(f, sheets, args.format), fh, indent=2, ensure_ascii=False)
        print(f"[OK] {f.name} -> {out}  ({total} entries)")

    if args.merge:
        res = {"source_files": [f.name for f in files], "format": args.format,
               "timetable": shape(merged, args.format)}
        with open(args.merge, "w", encoding="utf-8") as fh:
            json.dump(res, fh, indent=2, ensure_ascii=False)
        print(f"[OK] merged -> {args.merge}  ({len(merged)} entries)")


if __name__ == "__main__":
    main()
